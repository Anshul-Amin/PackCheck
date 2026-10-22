"""Approved content sheet ingestion module using openpyxl."""

import re
from pathlib import Path

import openpyxl

from src.packcheck.models import ApprovedData, NutritionData, NutritionItem


def _parse_amount_and_unit(val: str | float | int | None) -> tuple[float | None, str]:
    """Parse numeric amount and unit from strings like '15g', '89mcg', '166', '45mg'."""
    if val is None:
        return None, "g"
    if isinstance(val, (int, float)):
        return float(val), "g"

    s = str(val).strip().replace("–", "").replace("-", "")
    if not s or s.lower() == "n/a":
        return None, "g"

    match = re.search(r"([\d.]+)\s*([a-zA-Z%]+)?", s)
    if match:
        amt = float(match.group(1))
        unit = match.group(2) or "g"
        return amt, unit
    return None, "g"


def _parse_dv(val: str | float | int | None) -> float | None:
    """Parse %DV from values like '12%', '2.2%*', 12."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)

    s = str(val).strip().replace("*", "").replace("%", "").strip()
    if not s or s in ("–", "-") or s.lower() == "n/a":
        return None

    try:
        return float(s)
    except ValueError:
        return None


def _clean_ingredients_list(raw_text: str) -> list[str]:
    """Parse comma-separated ingredient list into cleaned items."""
    # Strip leading 'INGREDIENTS:' label if present
    text = re.sub(r"^\s*ingredients\s*:\s*", "", raw_text, flags=re.IGNORECASE).strip()
    if not text:
        return []

    # Split by comma (ignoring parenthetical commas for now or standard split)
    # E.g. 'Organic Wholegrain Oat Flakes (79%), Almonds, Pecans...'
    items = []
    # Regex split on commas outside parentheses
    raw_tokens = re.split(r",\s*(?![^()]*\))", text)
    for token in raw_tokens:
        clean = token.strip().rstrip(".")
        if clean:
            items.append(clean)
    return items


def ingest_content_sheet(xlsx_path: str | Path) -> ApprovedData:
    """Parse approved product information from the content sheet."""
    path = Path(xlsx_path)
    if not path.exists():
        raise FileNotFoundError(f"Content sheet not found at: {path}")

    wb = openpyxl.load_workbook(str(path), data_only=True)
    sheet_name = "Content Sheet - US" if "Content Sheet - US" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    approved = ApprovedData()
    notes: dict[str, str] = {}

    # Read row by row and map key fields
    for row in range(1, ws.max_row + 1):
        cell_a = ws.cell(row=row, column=1).value
        cell_b = ws.cell(row=row, column=2).value
        cell_c = ws.cell(row=row, column=3).value
        cell_d = ws.cell(row=row, column=4).value
        cell_e = ws.cell(row=row, column=5).value
        cell_f = ws.cell(row=row, column=6).value

        # Collect notes
        if cell_a and cell_c:
            notes[str(cell_a).strip()] = str(cell_c).strip()
        if cell_d and cell_f:
            notes[str(cell_d).strip()] = str(cell_f).strip()

        # Product Name
        if cell_d and "product name" in str(cell_d).lower():
            approved.product_name = str(cell_e).strip() if cell_e else None

        # Net Weight
        if cell_a and "portions selected" in str(cell_a).lower():
            approved.net_weight_raw = str(cell_b).strip() if cell_b else None
            amt, _ = _parse_amount_and_unit(cell_b)
            if amt:
                approved.net_weight_g = amt
        elif cell_a and "package size" in str(cell_a).lower():
            if not approved.net_weight_g:
                amt, _ = _parse_amount_and_unit(cell_b)
                if amt:
                    approved.net_weight_g = amt

        # UPC
        if cell_d and "upc" in str(cell_d).lower():
            approved.upc = str(cell_e).strip() if cell_e else None

        # Ingredients
        if cell_a and "ingredient list" in str(cell_a).lower() and cell_b:
            approved.raw_ingredients_text = str(cell_b).strip()
            approved.ingredients = _clean_ingredients_list(str(cell_b))

        # Allergens
        if cell_a and "allergen" in str(cell_a).lower() and cell_b:
            approved.raw_allergens_text = str(cell_b).strip()
            # Clean lines
            lines = str(cell_b).splitlines()
            for line in lines:
                cleaned_line = re.sub(
                    r"^\s*allergen\s+information\s*:\s*", "", line, flags=re.IGNORECASE
                ).strip()
                if cleaned_line:
                    for item in cleaned_line.split(","):
                        if item.strip():
                            approved.allergens.append(item.strip())

        # Claims & Certifications
        if cell_a and "claims" in str(cell_a).lower() and cell_b:
            raw_claims = str(cell_b).split(";")
            approved.claims = [c.strip() for c in raw_claims if c.strip()]
        elif cell_d and "notes" in str(cell_d).lower() and cell_e and not approved.claims:
            # Notes often carry claims e.g. "Organic certified (USDA Organic), Certified Vegan..."
            raw_claims = str(cell_e).split(",")
            approved.claims = [c.strip() for c in raw_claims if c.strip()]

        # Dietary information
        if cell_a and "dietary information" in str(cell_a).lower() and cell_b:
            approved.dietary_info = str(cell_b).strip()

        # Brand / Manufacturer
        if cell_a and "brand" in str(cell_a).lower() and cell_b:
            approved.manufacturer = str(cell_b).strip()

    # Parse Nutrition Table from row 46 onwards
    nutrition = NutritionData()

    # Check Servings per package (typically row 34)
    for row in range(30, 45):
        ca = ws.cell(row=row, column=1).value
        cb = ws.cell(row=row, column=2).value
        if ca and "servings per package" in str(ca).lower():
            amt, _ = _parse_amount_and_unit(cb)
            nutrition.servings_per_container = amt
        elif ca and "serving size" in str(ca).lower():
            nutrition.serving_size_household = str(cb).strip() if cb else None
            amt, _ = _parse_amount_and_unit(cb)
            if amt:
                nutrition.serving_size_g = amt

    # Parse Nutrients
    nutrient_key_map = {
        "calories": "calories",
        "total fat": "total_fat",
        "saturated fat": "saturated_fat",
        "trans fat": "trans_fat",
        "cholesterol": "cholesterol",
        "sodium": "sodium",
        "total carbohydrate": "total_carbohydrate",
        "dietary fiber": "dietary_fiber",
        "total sugars": "total_sugars",
        "incl. added sugars": "added_sugars",
        "added sugars": "added_sugars",
        "protein": "protein",
        "vitamin d": "vitamin_d",
        "calcium": "calcium",
        "iron": "iron",
        "potassium": "potassium",
    }

    for row in range(46, ws.max_row + 1):
        ca = ws.cell(row=row, column=1).value
        cb = ws.cell(row=row, column=2).value
        cc = ws.cell(row=row, column=3).value

        if not ca:
            continue

        raw_name = str(ca).strip().lower()
        matched_key = None
        for pattern, key in nutrient_key_map.items():
            if pattern in raw_name:
                matched_key = key
                break

        if not matched_key:
            continue

        amt, unit = _parse_amount_and_unit(cb)
        dv = _parse_dv(cc)

        if matched_key == "calories":
            nutrition.calories = amt
        else:
            nutrition.nutrients[matched_key] = NutritionItem(
                name=str(ca).strip(),
                amount=amt,
                unit=unit,
                dv_percent=dv,
            )

    approved.nutrition = nutrition
    approved.notes = notes
    wb.close()
    return approved
